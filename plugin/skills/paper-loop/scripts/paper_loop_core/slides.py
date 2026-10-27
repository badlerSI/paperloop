"""Slides: a small model of a deck, laid out here and written two ways.

A slide is a list of elements placed in points on a 960 x 540 canvas (13.333 x 7.5 inches, 16:9), with the
origin at the top left. `write_pptx` writes an editable PowerPoint file with python-pptx, which Claude's
sandboxes include; `write_pdf` draws the same slides with reportlab, for looking at them and as a copy to
present from. Text is measured with Helvetica's widths, which Arial (the font the .pptx names) shares, and
wrapped a little short of each box, so PowerPoint breaks lines where the PDF does or fits a word more.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from reportlab.lib.colors import HexColor
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from . import fonts as F

W, H = 960.0, 540.0
MARGIN = 48.0
EMU = 12700  # per point
SLACK = 0.96  # wrap at this share of a box's width
WHITE, INK, MUTED = "FFFFFF", "1F2328", "5F6670"


# --- The model ---------------------------------------------------------------------------

@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    color: str | None = None  # hex without "#"; None takes the box's color


@dataclass
class Para:
    runs: list[Run]
    bullet: str | None = None  # "•" for a bulleted point
    bullet_color: str | None = None
    space_after: float = 0.0


@dataclass
class Text:
    x: float
    y: float
    w: float
    h: float
    paras: list[Para]
    size: float
    color: str = INK
    align: str = "left"  # left, center, right
    valign: str = "top"  # top, middle
    mono: bool = False
    leading: float = 1.2
    role: str = ""  # "title" goes in the slide's title placeholder, so outlines and screen readers find it;
    # "answer" is a key's answer written onto a question's picture
    field: str = ""  # "slidenum": the slide's number, kept right when slides move


@dataclass
class Box:
    x: float
    y: float
    w: float
    h: float
    fill: str | None
    line: str | None = None
    line_w: float = 0.0
    shape: str = "round"  # round, rect, oval
    radius: float = 12.0


@dataclass
class Picture:
    x: float
    y: float
    w: float
    h: float
    path: str
    alt: str = ""
    border: str | None = None


@dataclass
class Bars:
    """A horizontal bar chart: a native chart in the .pptx."""
    x: float
    y: float
    w: float
    h: float
    labels: list[str]
    values: list[float]
    colors: list[str]
    size: float
    color: str = INK
    suffix: str = ""  # "%" after each value


@dataclass
class Slide:
    title: str
    elements: list = field(default_factory=list)
    notes: str = ""
    background: str = WHITE


_GLUE = re.compile(r"(?<=\S) ([:=×÷+→<>≤≥−-]) (?=\S)")


def runs(text: str, color: str | None = None, bold: bool = False, italic: bool = False,
         markup: bool = True) -> list[Run]:
    """Text with **bold** marked, as runs. "2 : 3" and "3 × 4" are kept on one line (no-break spaces).
    markup=False for text from a worksheet or a student, where ** is Python's power, not bold."""
    text = _GLUE.sub("\u00a0\\1\u00a0", text)
    if not markup:
        return [Run(text, bold=bold, italic=italic, color=color)]
    out = []
    for k, piece in enumerate(re.split(r"\*\*", text)):
        if piece:
            out.append(Run(piece, bold=bold or k % 2 == 1, italic=italic, color=color))
    return out or [Run("", bold=bold, italic=italic, color=color)]


def plain(text: str) -> str:
    return text.replace("**", "")


# --- Measuring text ----------------------------------------------------------------------

class Measure:
    """Widths and line breaks, in the fonts the PDF draws with: Helvetica, which has Arial's widths, and
    DejaVu Sans for the few characters Helvetica lacks (π, √, →), as PowerPoint borrows another font for them."""

    def __init__(self, texts=()):
        self.u = F.unicode_fonts()  # None when no Unicode font is installed: those characters become "?"
        self.missing: set[str] = set()  # characters no font here can draw (Chinese, emoji): blank in the PDF
        self._glyphs = None

    def font(self, bold: bool, italic: bool, mono: bool) -> str:
        base = "Courier" if mono else "Helvetica"
        style = ("Bold" if bold else "") + ("Oblique" if italic else "")
        return f"{base}-{style}" if style else base

    def _fallback(self, bold: bool, italic: bool, mono: bool) -> str:
        return self.u.mono if mono else self.u.bold if bold else self.u.oblique if italic else self.u.regular

    def segments(self, s: str, bold: bool = False, italic: bool = False, mono: bool = False) -> list[tuple[str, str]]:
        """s split into (text, font) pieces, so each character is drawn by a font that has it."""
        out: list[tuple[str, str]] = []
        for ch in s:
            try:
                ch.encode("cp1252")
                piece, font = ch, self.font(bold, italic, mono)
            except UnicodeEncodeError:
                if self.u:
                    piece, font = ch, self._fallback(bold, italic, mono)
                    if not self._has(ch) and not ch.isspace():
                        self.missing.add(ch)
                else:
                    piece, font = F.HELVETICA.text(ch), self.font(bold, italic, mono)
                    self.missing.add(ch)
            if out and out[-1][1] == font:
                out[-1] = (out[-1][0] + piece, font)
            else:
                out.append((piece, font))
        return out

    def _has(self, ch: str) -> bool:
        if self._glyphs is None:
            from reportlab.pdfbase.pdfmetrics import getFont
            self._glyphs = set(getFont(self.u.regular).face.charToGlyph)
        return ord(ch) in self._glyphs

    def text(self, s: str) -> str:
        return s if self.u else F.HELVETICA.text(s)

    def width(self, s: str, run: Run, size: float, mono: bool = False) -> float:
        return sum(stringWidth(t, f, size) for t, f in self.segments(s, run.bold, run.italic, mono))

    def draw(self, c, x: float, y: float, s: str, size: float, bold: bool = False, italic: bool = False,
             mono: bool = False) -> float:
        """Draws s from x on baseline y; returns where it ended."""
        for t, f in self.segments(s, bold, italic, mono):
            c.setFont(f, size)
            c.drawString(x, y, t)
            x += stringWidth(t, f, size)
        return x

    def indent(self, para: Para, size: float) -> float:
        return size * 0.95 if para.bullet else 0.0

    def wrap(self, para: Para, width: float, size: float, mono: bool = False) -> list[list[tuple[str, Run]]]:
        width = width * SLACK - self.indent(para, size)
        lines, line, line_w = [], [], 0.0
        fresh = True  # at the start of the paragraph or after a newline: leading spaces are indentation

        def close():
            nonlocal line, line_w
            while line and line[-1][0].isspace() and len(line) > 1:
                line.pop()
            lines.append(line)
            line, line_w = [], 0.0

        for run in para.runs:
            for piece in re.split(r"(\n| +)", run.text):
                if not piece:
                    continue
                if piece == "\n":
                    close()
                    fresh = True
                    continue
                w = self.width(piece, run, size, mono)
                if piece.isspace():
                    if line or fresh:
                        line.append((piece, run))
                        line_w += w
                    continue
                fresh = False
                if line and line_w + w > width:
                    close()
                while w > width and len(piece) > 1:  # a word longer than the line
                    cut = len(piece) - 1
                    while cut > 1 and self.width(piece[:cut], run, size, mono) > width:
                        cut -= 1
                    line.append((piece[:cut], run))
                    close()
                    piece = piece[cut:]
                    w = self.width(piece, run, size, mono)
                line.append((piece, run))
                line_w += w
        close()
        return lines

    def height(self, paras: list[Para], width: float, size: float, leading: float = 1.2, mono: bool = False) -> float:
        lines = sum(len(self.wrap(p, width, size, mono)) for p in paras)
        return lines * size * leading + sum(p.space_after for p in paras[:-1])

    def line_width(self, line, size: float, mono: bool = False) -> float:
        return sum(self.width(s, r, size, mono) for s, r in line)

    def fit(self, paras: list[Para], width: float, height: float, size: float, floor: float,
            leading: float = 1.2, mono: bool = False) -> tuple[float, bool]:
        """The largest size from `size` down to `floor` at which the text fits; (floor, False) if none does."""
        s = size
        while s >= floor:
            scaled = [Para(p.runs, p.bullet, p.bullet_color, p.space_after * s / size) for p in paras]
            if self.height(scaled, width, s, leading, mono) <= height + 0.5:
                return s, True
            s -= 1
        return floor, False


# --- Colors ------------------------------------------------------------------------------

def mix(color: str, other: str, share: float) -> str:
    """color moved `share` of the way to other."""
    a = [int(color[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(other[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * share):02X}" for x, y in zip(a, b))


@dataclass(frozen=True)
class Theme:
    accent: str
    dark: str

    @property
    def tint(self) -> str:  # card backgrounds
        return mix(self.accent, WHITE, 0.9)

    @property
    def soft(self) -> str:  # chips and quiet fills
        return mix(self.accent, WHITE, 0.75)

    @property
    def on_dark(self) -> str:  # subtitles on a dark slide
        return mix(self.accent, WHITE, 0.6)


# Whole words in the subject pick the deck's colors. Language arts comes before world languages, physical
# education before physical science, and computer science before science.
_SUBJECT_COLORS = (
    (r"english|ela|reading|writing|literature|language arts|composition|poetry|literacy", "9C2F5A", "3A1426"),
    (r"spanish|french|german|chinese|mandarin|japanese|latin|italian|arabic|korean|asl|world languages?|esl|ell",
     "C2412D", "3D1510"),
    (r"health|pe|physical education|gym", "B86E00", "3D2600"),
    (r"math|mathematics|algebra|geometry|calculus|statistics|precalculus|trigonometry|arithmetic", "2453A6",
     "13254F"),
    (r"computer|computing|coding|programming|cs|technology|robotics", "00796B", "0D2F2B"),
    (r"science|biology|chemistry|physics|earth|environmental|anatomy|ecology|astronomy|stem", "2E7D4F", "143A26"),
    (r"history|social studies|civics|government|geography|economics|psychology|sociology", "B2532A", "3F1D10"),
    (r"music|band|choir|orchestra|art|arts|drama|theater|theatre|dance", "6B3FA0", "2A1845"),
)


def theme_for(subject: str) -> Theme:
    s = subject.lower()
    for pattern, accent, dark in _SUBJECT_COLORS:
        if re.search(rf"\b(?:{pattern})\b", s):
            return Theme(accent, dark)
    return Theme("D0661C", "263238")


def sizes_for(grade: int) -> dict:
    """Type sizes for a classroom projector: the back row has to read it, so bigger than a meeting deck."""
    if grade <= 3:
        return dict(title=44, body=32, small=24, floor=24, dense=20, words=45)
    if grade <= 8:
        return dict(title=40, body=28, small=22, floor=20, dense=17, words=70)
    return dict(title=36, body=26, small=20, floor=18, dense=16, words=90)


# --- PowerPoint --------------------------------------------------------------------------

def write_pptx(slides: list[Slide], path: Path, title: str) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(W * EMU)), Emu(int(H * EMU))
    prs.core_properties.title = title
    prs.core_properties.author = "Paper Loop"
    layout = prs.slide_layouts[5]  # Title Only
    for n, s in enumerate(slides, start=1):
        slide = prs.slides.add_slide(layout)
        if s.background != WHITE:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(s.background)
        placeholder = slide.shapes.title
        has_title = False
        for e in s.elements:
            if isinstance(e, Box):
                _pptx_box(slide, e)
            elif isinstance(e, Picture):
                _pptx_picture(slide, e)
            elif isinstance(e, Bars):
                _pptx_bars(slide, e)
            elif isinstance(e, Text) and e.role == "title" and not has_title:
                _place(placeholder, e)
                _fill_frame(placeholder.text_frame, e, n)
                has_title = True
            elif isinstance(e, Text):
                shape = slide.shapes.add_textbox(*_emu(e))
                _fill_frame(shape.text_frame, e, n)
        if not has_title:
            placeholder._element.getparent().remove(placeholder._element)
        if s.notes:
            slide.notes_slide.notes_text_frame.text = s.notes
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


def _rgb(color: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(color)


def _emu(e) -> tuple:
    from pptx.util import Emu
    return tuple(Emu(int(round(v * EMU))) for v in (e.x, e.y, e.w, e.h))


def _place(shape, e) -> None:
    shape.left, shape.top, shape.width, shape.height = _emu(e)


def _fill_frame(tf, e: Text, slide_no: int) -> None:
    from lxml import etree
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_top = tf.margin_bottom = 0
    tf.margin_right = int(e.w * (1 - SLACK) * EMU)  # PowerPoint wraps where the PDF (and our measuring) wraps
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if e.valign == "middle" else MSO_ANCHOR.TOP
    align = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[e.align]
    font_name = "Courier New" if e.mono else "Arial"
    for i, para in enumerate(e.paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = Pt(e.size * e.leading)
        if para.space_after:
            p.space_after = Pt(para.space_after)
        if para.bullet:
            pPr = p._p.get_or_add_pPr()
            indent = int(e.size * 0.95 * EMU)
            pPr.set("marL", str(indent))
            pPr.set("indent", str(-indent))
            clr = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr")).set("val", para.bullet_color or e.color)
            etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", para.bullet)
        if e.field == "slidenum":
            fld = etree.SubElement(p._p, qn("a:fld"))
            fld.set("id", "{6F2A3B9C-1D4E-4F5A-8B7C-2E3D4F5A6B7C}")
            fld.set("type", "slidenum")
            rpr = etree.SubElement(fld, qn("a:rPr"))
            rpr.set("lang", "en-US")
            rpr.set("sz", str(int(e.size * 100)))
            fill = etree.SubElement(rpr, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", e.color)
            etree.SubElement(rpr, qn("a:latin")).set("typeface", font_name)
            etree.SubElement(fld, qn("a:t")).text = str(slide_no)
            continue
        for run in para.runs:
            r = p.add_run()
            r.text = run.text
            f = r.font
            f.size = Pt(e.size)
            f.bold = run.bold
            f.italic = run.italic
            f.name = font_name
            f.color.rgb = _rgb(run.color or e.color)


def _pptx_box(slide, e: Box) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt

    kind = {"round": MSO_SHAPE.ROUNDED_RECTANGLE, "rect": MSO_SHAPE.RECTANGLE, "oval": MSO_SHAPE.OVAL}[e.shape]
    shape = slide.shapes.add_shape(kind, *_emu(e))
    if e.shape == "round":
        shape.adjustments[0] = min(0.5, e.radius / max(1.0, min(e.w, e.h)))
    if e.fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(e.fill)
    else:
        shape.fill.background()
    if e.line:
        shape.line.color.rgb = _rgb(e.line)
        shape.line.width = Pt(e.line_w or 1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False


def _pptx_picture(slide, e: Picture) -> None:
    from pptx.util import Pt

    pic = slide.shapes.add_picture(e.path, *_emu(e))
    if e.alt:
        pic._element.nvPicPr.cNvPr.set("descr", e.alt)
    if e.border:
        pic.line.color.rgb = _rgb(e.border)
        pic.line.width = Pt(1)


def _pptx_bars(slide, e: Bars) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Pt

    data = CategoryChartData()  # a bar chart draws its first category at the bottom: feed it bottom-up
    data.categories = e.labels[::-1]
    data.add_series("Students", e.values[::-1])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, *_emu(e), data).chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = "Arial"
    chart.font.size = Pt(e.size)
    chart.font.color.rgb = _rgb(e.color)
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.vary_by_categories = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = '0"' + e.suffix + '"' if e.suffix else "0"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(e.size)
    labels.font.bold = True
    series = plot.series[0]
    for i, color in enumerate(e.colors[::-1]):
        point = series.points[i]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = _rgb(color)
    cats = chart.category_axis
    cats.has_major_gridlines = False
    cats.tick_labels.font.size = Pt(e.size)
    cats.format.line.fill.background()
    values = chart.value_axis
    values.visible = False
    values.has_major_gridlines = False
    values.minimum_scale = 0
    values.maximum_scale = max([*e.values, 1]) * 1.18


# --- PDF ---------------------------------------------------------------------------------

def write_pdf(slides: list[Slide], path: Path, title: str, m: Measure) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=(W, H))
    c.setTitle(m.text(title))
    c.setAuthor("Paper Loop")
    for n, s in enumerate(slides, start=1):
        c.setFillColor(HexColor("#" + s.background))
        c.rect(0, 0, W, H, stroke=0, fill=1)
        for e in s.elements:
            if isinstance(e, Box):
                _pdf_box(c, e)
            elif isinstance(e, Picture):
                c.drawImage(e.path, e.x, H - e.y - e.h, e.w, e.h)
                if e.border:
                    c.setStrokeColor(HexColor("#" + e.border))
                    c.setLineWidth(1)
                    c.rect(e.x, H - e.y - e.h, e.w, e.h, stroke=1, fill=0)
            elif isinstance(e, Bars):
                _pdf_bars(c, e, m)
            elif isinstance(e, Text):
                _pdf_text(c, e, m, n)
        c.showPage()
    c.save()


def _pdf_box(c: canvas.Canvas, e: Box) -> None:
    c.setFillColor(HexColor("#" + (e.fill or WHITE)))
    if e.line:
        c.setStrokeColor(HexColor("#" + e.line))
        c.setLineWidth(e.line_w or 1)
    stroke, fill = (1 if e.line else 0), (1 if e.fill else 0)
    y = H - e.y - e.h
    if e.shape == "oval":
        c.ellipse(e.x, y, e.x + e.w, y + e.h, stroke=stroke, fill=fill)
    elif e.shape == "round":
        c.roundRect(e.x, y, e.w, e.h, min(e.radius, e.w / 2, e.h / 2), stroke=stroke, fill=fill)
    else:
        c.rect(e.x, y, e.w, e.h, stroke=stroke, fill=fill)


def _pdf_text(c: canvas.Canvas, e: Text, m: Measure, slide_no: int) -> None:
    if e.field == "slidenum":
        e = Text(e.x, e.y, e.w, e.h, [Para([Run(str(slide_no))])], e.size, e.color, e.align, e.valign)
    step = e.size * e.leading
    laid = [(p, m.wrap(p, e.w, e.size, e.mono)) for p in e.paras]
    total = sum(len(lines) for _, lines in laid) * step + sum(p.space_after for p in e.paras[:-1])
    top = e.y + ((e.h - total) / 2 if e.valign == "middle" else 0)
    y = top
    for p, lines in laid:
        indent = m.indent(p, e.size)
        for k, line in enumerate(lines):
            baseline = H - (y + (step - e.size) / 2 + e.size * 0.8)
            width = m.line_width(line, e.size, e.mono) + indent
            room = e.w * SLACK  # as in the .pptx, whose right inset is the rest
            x = e.x + {"left": 0, "center": (room - width) / 2, "right": room - width}[e.align]
            if p.bullet and k == 0:
                c.setFillColor(HexColor("#" + (p.bullet_color or e.color)))
                m.draw(c, x, baseline, p.bullet, e.size)
            x += indent
            for piece, run in line:
                c.setFillColor(HexColor("#" + (run.color or e.color)))
                x = m.draw(c, x, baseline, piece, e.size, run.bold, run.italic, e.mono)
            y += step
        y += p.space_after


def _pdf_bars(c: canvas.Canvas, e: Bars, m: Measure) -> None:
    n = max(1, len(e.labels))
    label_w = max(m.width(lbl, Run(lbl), e.size) for lbl in e.labels) + 14 if e.labels else 0
    value_w = m.width("888" + e.suffix, Run("", bold=True), e.size) + 10
    span = max([*e.values, 1]) * 1.0
    row = e.h / n
    bar_h = row / (1 + 0.45)
    for i, (label, value, color) in enumerate(zip(e.labels, e.values, e.colors)):
        top = e.y + i * row + (row - bar_h) / 2
        mid = H - (top + bar_h / 2)
        c.setFillColor(HexColor("#" + e.color))
        m.draw(c, e.x + label_w - 10 - m.width(label, Run(label), e.size), mid - e.size * 0.35, label, e.size)
        length = (e.w - label_w - value_w) * value / span
        c.setFillColor(HexColor("#" + color))
        c.rect(e.x + label_w, H - top - bar_h, max(length, 0.5), bar_h, stroke=0, fill=1)
        c.setFillColor(HexColor("#" + e.color))
        m.draw(c, e.x + label_w + length + 6, mid - e.size * 0.35, f"{value:g}{e.suffix}", e.size, bold=True)


# --- Checks ------------------------------------------------------------------------------

def out_of_bounds(s: Slide) -> list[str]:
    out = []
    for e in s.elements:
        if e.x < -0.5 or e.y < -0.5 or e.x + e.w > W + 0.5 or e.y + e.h > H + 0.5:
            out.append(f"{type(e).__name__} at ({e.x:.0f}, {e.y:.0f}, {e.w:.0f}x{e.h:.0f}) runs off the slide")
    return out


def words_on(s: Slide) -> int:
    return sum(len(plain(r.text).split()) for e in s.elements if isinstance(e, Text) and e.role not in ("title", "answer")
               and not e.field for p in e.paras for r in p.runs)
